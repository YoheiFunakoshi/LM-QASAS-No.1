"""Read the provided Human IGH report layout without editing the workbook.

Back_data G:Q is the complete record block, starting on row 1 without headers.
T:AA is a top-50 ranking and must never be used as the repertoire input.
"""
from __future__ import annotations

from collections import Counter
from contextlib import contextmanager
import io
import math
from pathlib import Path
import re
from zipfile import ZipFile, BadZipFile

from .inputs import InputValidationError, _gene_set, _isotype, _CANONICAL_AA

FORMAT = 'takara_rg_xlsx'
POLICY = 'takara-rg-hIGH20181210-v1'
FIELDS = ('V', 'V_function', 'D', 'D_function', 'J', 'J_function',
          'C', 'C_function', 'CDR3', 'frame', 'count')
MAX_EXPANDED_BYTES = 256 * 1024 * 1024
MAX_ROWS = 250000


def _integer(value):
    return (type(value) in (int, float) and math.isfinite(value)
            and value >= 0 and int(value) == value)


@contextmanager
def _open(data: bytes):
    # Inspect the archive in memory. Never extract, refresh links, or save it.
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if (len(entries) > 4096 or sum(e.file_size for e in entries) > MAX_EXPANDED_BYTES
                    or len({e.filename for e in entries}) != len(entries)
                    or any('vbaproject' in e.filename.lower() for e in entries)):
                raise InputValidationError('Excel archive exceeds supported limits or contains macros.')
            required = {'[Content_Types].xml', 'xl/workbook.xml'}
            if not required <= {e.filename for e in entries}:
                raise InputValidationError('Not a supported XLSX workbook.')
    except (BadZipFile, OSError, ValueError, RuntimeError):
        raise InputValidationError('Excel file is not a readable XLSX archive.') from None
    from openpyxl import load_workbook
    from openpyxl.xml import DEFUSEDXML
    if not DEFUSEDXML:
        raise InputValidationError('Excel XML protection requires defusedxml in the project environment.')
    book = None
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        yield book
    except InputValidationError:
        raise
    except Exception:
        raise InputValidationError('Excel workbook could not be read; check the supported report layout.') from None
    finally:
        if book is not None:
            book.close()


def _inspect(book):
    if book.sheetnames[:2] != ['PRINT_hIGH', 'Back_data']:
        raise InputValidationError('Excel requires PRINT_hIGH followed by Back_data as the second sheet.')
    front, sheet = book.worksheets[:2]
    if (front['F1'].value != 'Repertoire Analysis Report (Human IGH)'
            or not isinstance(front['F5'].value, str)
            or not front['F5'].value.endswith('Sheet ver. hIGH20181210')):
        raise InputValidationError('Excel Human IGH report version is not supported.')
    declared_rows = sheet.max_row
    # Stored dimensions may understate the complete data block.
    sheet.reset_dimensions()
    top = list(sheet.iter_rows(min_row=1, max_row=7, min_col=1, max_col=19))
    anchors = {(1, 6): '7:All Data', (1, 19): '8:In frame Data', (2, 19): 'Ranking',
               (4, 2): 'Total reads', (5, 2): 'Assigned reads',
               (6, 2): 'In frame', (7, 2): 'Unique reads (In frame)'}
    if any(top[r-1][c-1].value != value for (r, c), value in anchors.items()):
        raise InputValidationError('Excel Back_data layout markers do not match the supported report.')
    summary = {name: top[r-1][2].value for r, name in
               ((4, 'total_reads'), (5, 'assigned_reads'), (6, 'in_frame_reads'), (7, 'in_frame_unique'))}
    if any(not _integer(value) for value in summary.values()):
        raise InputValidationError('Excel summary counts must be stored nonnegative integers, not formulas.')
    if not summary['total_reads'] >= summary['assigned_reads'] >= summary['in_frame_reads']:
        raise InputValidationError('Excel summary counts are inconsistent.')
    return sheet, {'input_format': FORMAT, 'input_policy': POLICY, 'source_sheet': sheet.title,
                   'source_sheet_index': 2, 'source_columns': 'G:Q', 'source_data_start_row': 1,
                   'reported_max_row': declared_rows, 'dimensions_reset': True,
                   'isotype_granularity': 'subclass',
                   'cdr3_definition': 'as_reported_no_boundary_repair',
                   'summary': summary}


def check_workbook(data: bytes) -> dict:
    """Preflight only: signature, second worksheet and summary cell types."""
    with _open(data) as book:
        _, audit = _inspect(book)
        return audit


def _calls(value, prefix):
    if not isinstance(value, str):
        return None
    # This report uses comma-separated alternatives; slashes inside gene names stay intact.
    if '//' in value:
        return None
    if prefix == 'IGHD':
        # This report also uses lowercase a/b suffixes for D gene copies.
        # Validate them without changing the original annotation or its case.
        candidates = [item.strip() for item in value.split(',')]
        pattern = r'IGHD[A-Z0-9][A-Z0-9/-]*[ab]?(?:\*[A-Z0-9]+)?\Z'
        if not all(re.fullmatch(pattern, item) for item in candidates):
            return None
        return tuple(sorted({item.split('*', 1)[0] for item in candidates}))
    return _gene_set('//'.join(value.split(',')), prefix)


def _constant(value):
    if not isinstance(value, str) or '//' in value:
        return None
    candidates = [_isotype(item) for item in value.split(',')]
    if any(item is None for item in candidates) or len(set(candidates)) != 1:
        return None
    return candidates[0]


def _safe_raw(value):
    # Preserve ordinary typed source cells; reject rather than invent strings for other types.
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float) and math.isfinite(value):
        return value
    raise InputValidationError('Excel All Data contains an unsupported cell type.')


def parse_workbook(data: bytes, path: Path, timepoint: str, subject: str):
    with _open(data) as book:
        sheet, audit = _inspect(book)
        grouped, exclusions, reasons_count = {}, [], Counter()
        total = accepted = ambiguous_v = ambiguous_j = last_data_row = 0
        assigned_reads = in_frame_reads = 0
        in_frame_keys = set()
        for row_number, cells in enumerate(sheet.iter_rows(min_row=1, min_col=7, max_col=17), 1):
            if row_number > MAX_ROWS:
                raise InputValidationError('Excel exceeds the supported worksheet row limit.')
            if all(cell.value is None for cell in cells):
                continue
            if any(cell.data_type in ('f', 'e') for cell in cells):
                raise InputValidationError('Excel All Data must contain stored values, not formulas or errors.')
            raw = dict(zip(FIELDS, (_safe_raw(cell.value) for cell in cells), strict=True))
            total += 1
            last_data_row = row_number
            if not _integer(raw['count']):
                raise InputValidationError('Excel All Data counts must be nonnegative integers.')
            assigned_reads += int(raw['count'])
            if raw['frame'] == 'in-frame':
                in_frame_reads += int(raw['count'])
                in_frame_keys.add(tuple(raw[key] for key in ('V', 'D', 'J', 'CDR3', 'C')))
            reasons = []
            cdr3 = raw['CDR3']
            if not isinstance(cdr3, str) or not _CANONICAL_AA.fullmatch(cdr3):
                reasons.append('cdr3_noncanonical')
            if not isinstance(cdr3, str) or len(cdr3) < 5:
                reasons.append('cdr3_too_short')
            if raw['frame'] != 'in-frame':
                reasons.append('frame_not_in_frame')
            for field in ('V', 'D', 'J'):
                if raw[field + '_function'] != 'F':
                    reasons.append(field.lower() + '_function_not_F')
            v_genes, j_genes = _calls(raw['V'], 'IGHV'), _calls(raw['J'], 'IGHJ')
            if v_genes is None:
                reasons.append('v_annotation_invalid')
            if j_genes is None:
                reasons.append('j_annotation_invalid')
            if _calls(raw['D'], 'IGHD') is None:
                reasons.append('d_annotation_invalid')
            isotype = _constant(raw['C'])
            if isotype is None:
                reasons.append('cseg_unmapped_or_ambiguous')
            if reasons:
                exclusions.append({'source_row': row_number, 'source_sheet': sheet.title, 'reasons': reasons})
                reasons_count.update(reasons)
                continue
            accepted += 1
            ambiguous_v += len(v_genes) > 1
            ambiguous_j += len(j_genes) > 1
            key = (v_genes, j_genes, cdr3, isotype)
            if key not in grouped:
                grouped[key] = {'cdr3': cdr3, 'v_gene': '//'.join(v_genes), 'j_gene': '//'.join(j_genes),
                                'isotype': isotype, 'timepoint': timepoint, 'subject': subject,
                                'source_file': str(path), 'source_sheet': sheet.title, 'source_columns': 'G:Q',
                                'source_row': row_number, 'source_rows': [], 'raw_records': []}
            grouped[key]['source_rows'].append(row_number)
            grouped[key]['raw_records'].append(raw)
        summary = audit['summary']
        if (assigned_reads != summary['assigned_reads'] or in_frame_reads != summary['in_frame_reads']
                or len(in_frame_keys) != summary['in_frame_unique']):
            raise InputValidationError('Excel All Data does not reconcile with report summary; input rejected.')
        if not grouped:
            raise InputValidationError(f'{timepoint}: no rows remain after input validation.')
        audit.update(source_file=str(path), total_rows=total, accepted_rows=accepted,
                     excluded_rows=len(exclusions), clone_count=len(grouped), merged_rows=accepted-len(grouped),
                     reason_counts=dict(sorted(reasons_count.items())), exclusions=exclusions,
                     rows_with_multiple_v_genes=ambiguous_v, rows_with_multiple_j_genes=ambiguous_j,
                     last_data_row=last_data_row, report_summary_reconciled=True,
                     report_functional_vdj_labels_required=['F'], report_frame_required='in-frame',
                     full_vdj_functionality_verified=False, nt_length_verified=False,
                     vendor_function_labels_independently_verified=False,
                     cdr3_boundaries_modified=False, counts_used_as_weights=False)
        return list(grouped.values()), audit
